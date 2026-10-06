from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    if level == 1:
        for run in h.runs:
            run.font.color.rgb = RGBColor(0, 102, 204)

def add_code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(10)
    p.style = 'No Spacing'
    p.paragraph_format.left_indent = Pt(20)

add_heading('Guía Definitiva de Conexión Frontend-Backend', 0)
doc.add_paragraph('Proyecto: Kawsay Eco-Dash')
doc.add_paragraph('Autor: Marcelo Zambrano (Backend)')
doc.add_paragraph('Destinatarios: Jordan y Narváez (Frontend)')

add_heading('1. Visión General')
doc.add_paragraph('¡Hola equipo! El servidor backend de Kawsay Eco-Dash ya está completamente funcional y desplegado en la nube. '
                  'Ya NO necesitan instalar Python, ni clonar el repositorio del backend, ni levantar servidores locales en sus computadoras.')
p = doc.add_paragraph('A partir de ahora, toda la comunicación se realizará haciendo peticiones HTTP (Fetch/Axios) a la siguiente ')
run = p.add_run('URL Base Pública:')
run.bold = True
add_code('https://kawsay-api.onrender.com')

add_heading('2. Documentación Interactiva (Swagger UI)', 2)
doc.add_paragraph('Pueden ver y probar visualmente todos los endpoints antes de programarlos en Next.js/React entrando desde su navegador a:')
add_code('https://kawsay-api.onrender.com/docs')

add_heading('3. Endpoints Disponibles y Ejemplos de Uso', 1)

add_heading('A. Obtener Botones de Hábitos', 2)
doc.add_paragraph('Este endpoint les devuelve la lista de acciones ambientales activas en la base de datos para que rendericen los botones en la interfaz.')
doc.add_paragraph('- Método: GET')
doc.add_paragraph('- Ruta: /api/v1/acciones/')
doc.add_paragraph('Ejemplo de Fetch en JavaScript:')
add_code('const obtenerAcciones = async () => {\n'
         '    const res = await fetch("https://kawsay-api.onrender.com/api/v1/acciones/");\n'
         '    const json = await res.json();\n'
         '    console.log(json.data); // Usen esto para hacer el .map() de los botones\n'
         '};')

add_heading('B. Registrar un Hábito (Click en un botón)', 2)
doc.add_paragraph('Al hacer clic en un botón, envían el UUID del usuario logueado en Supabase y el ID de la acción. El backend calculará el impacto, actualizará la racha, la XP y guardará todo en la base de datos. Además les devolverá un dato curioso (ODS 4) al azar.')
doc.add_paragraph('- Método: POST')
doc.add_paragraph('- Ruta: /api/v1/acciones/registrar')
doc.add_paragraph('- Body (JSON): { "usuario_id": "uuid-del-usuario", "accion_id": 1 }')
doc.add_paragraph('Ejemplo de Fetch en JavaScript:')
add_code('const registrarHabito = async (userId, actionId) => {\n'
         '    const res = await fetch("https://kawsay-api.onrender.com/api/v1/acciones/registrar", {\n'
         '        method: "POST",\n'
         '        headers: { "Content-Type": "application/json" },\n'
         '        body: JSON.stringify({ usuario_id: userId, accion_id: actionId })\n'
         '    });\n'
         '    const respuesta = await res.json();\n'
         '    \n'
         '    // Revisar si subió de nivel para mostrar animación:\n'
         '    if (respuesta.estado_usuario.subio_de_nivel) {\n'
         '        alert(¡Felicidades! Eres nivel );\n'
         '    }\n'
         '    \n'
         '    // Mostrar curiosidad educativa si la hay:\n'
         '    if (respuesta.recompensa_educativa) {\n'
         '        console.log("Sabías que:", respuesta.recompensa_educativa.contenido);\n'
         '    }\n'
         '};')

add_heading('C. Cargar el Dashboard Principal', 2)
doc.add_paragraph('Usen este endpoint al cargar la página principal del dashboard para pintar los gráficos (Chart.js / Recharts) y los totales numéricos (CO2 y Calorías).')
doc.add_paragraph('- Método: GET')
doc.add_paragraph('- Ruta: /api/v1/dashboard/{usuario_id}')
doc.add_paragraph('Ejemplo de Fetch en JavaScript:')
add_code('const cargarDashboard = async (userId) => {\n'
         '    const res = await fetch(https://kawsay-api.onrender.com/api/v1/dashboard/);\n'
         '    const dashboardData = await res.json();\n'
         '    \n'
         '    // Imprimir totales en pantalla\n'
         '    console.log("Total CO2 evitado:", dashboardData.impacto_global_co2);\n'
         '    console.log("Nivel:", dashboardData.nivel_actual, "Racha:", dashboardData.racha_dias);\n'
         '    \n'
         '    // Usar esto para el gráfico de Anillo (Donut Chart)\n'
         '    console.log("Data para Gráfico:", dashboardData.desglose_grafico);\n'
         '};')

add_heading('4. Consideraciones Finales', 1)
doc.add_paragraph('1. Autenticación: Recuerden que ustedes (Frontend) manejan el login directo con Supabase Auth. El ID que Supabase les entrega al hacer login exitoso, es el UUID que me deben pasar en las peticiones de los endpoints (usuario_id).')
doc.add_paragraph('2. Errores Comunes: Si reciben un Error 422 (Unprocessable Entity), significa que enviaron el JSON mal formado o les faltó el usuario_id.')

doc.save('Guia_Conexion_Frontend_Kawsay.docx')
print("Documento guardado exitosamente.")
